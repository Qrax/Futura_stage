import os
import calendar
import time
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter

# Mapping van Nederlandse maandnamen naar maandnummers
NEDERLANDS_MAANDEN = {
    "januari": 1, "februari": 2, "maart": 3, "april": 4,
    "mei": 5, "juni": 6, "juli": 7, "augustus": 8,
    "september": 9, "oktober": 10, "november": 11, "december": 12
}

def create_logbook_excel(nl_month: str, year: int, base_folder: str = "."):
    """
    Creëert een Excel logboek voor een specifieke maand en jaar.

    Args:
        nl_month: Nederlandse maandnaam (bijv. "augustus")
        year: Jaar (bijv. 2025)
        base_folder: Basismap waar het bestand wordt opgeslagen

    Returns:
        str: Pad naar het aangemaakte Excel bestand
    """
    month = nl_month.lower()
    if month not in NEDERLANDS_MAANDEN:
        raise ValueError(f"Maand '{nl_month}' is niet herkend. Gebruik een Nederlandse maandnaam.")

    month_number = NEDERLANDS_MAANDEN[month]

    # Bereken alle weekdagen van de maand
    start_date = datetime(year, month_number, 1)
    last_day = calendar.monthrange(year, month_number)[1]
    end_date = datetime(year, month_number, last_day)

    # Filter alleen weekdagen (maandag=0, vrijdag=4)
    weekdays = []
    current_date = start_date
    while current_date <= end_date:
        if current_date.weekday() < 5:  # 0-4 zijn weekdagen
            weekdays.append(current_date)
        current_date += timedelta(days=1)

    # Maak directory structuur aan
    year_folder = os.path.join(base_folder, str(year))
    os.makedirs(year_folder, exist_ok=True)

    # Bestandsnaam en pad
    filename = f"Logboek_Quincy_{year}_{nl_month.capitalize()}.xlsx"
    filepath = os.path.join(year_folder, filename)

    # Controleer of bestand al bestaat en open is
    if os.path.exists(filepath):
        try:
            os.rename(filepath, filepath)
        except OSError:
            timestamp = datetime.now().strftime("%H%M%S")
            filename = f"Logboek_Quincy_{year}_{nl_month.capitalize()}_{timestamp}.xlsx"
            filepath = os.path.join(year_folder, filename)
            print(f"⚠️  Origineel bestand is open. Nieuw bestand wordt aangemaakt: {filename}")

    # Maak workbook aan
    wb = Workbook()
    ws = wb.active
    ws.title = f"{nl_month.capitalize()} {year}"

    # Maak mooie header box
    header_end_row = _create_header_box(ws, nl_month.capitalize(), year)

    # Stel kolom headers in
    table_start_row = header_end_row + 2
    _setup_table_headers(ws, table_start_row)

    # Vul data in
    data_start_row = table_start_row + 1
    last_row = _fill_weekday_data(ws, weekdays, data_start_row)

    # Update formules met juiste bereik
    _update_summary_formulas(ws, data_start_row, last_row)

    # Stel kolom breedtes in
    _setup_column_widths(ws)

    # Bewaar het bestand met retry mechanisme
    max_retries = 3
    for attempt in range(max_retries):
        try:
            wb.save(filepath)
            break
        except PermissionError as e:
            if attempt < max_retries - 1:
                print(f"⚠️  Poging {attempt + 1} gefaald. Wacht 2 seconden...")
                time.sleep(2)
                timestamp = datetime.now().strftime("%H%M%S")
                filename = f"Logboek_Quincy_{year}_{nl_month.capitalize()}_{timestamp}.xlsx"
                filepath = os.path.join(year_folder, filename)
            else:
                raise PermissionError(
                    f"Kan bestand niet opslaan na {max_retries} pogingen. "
                    f"Controleer of:\n"
                    f"1. Het bestand niet open is in Excel\n"
                    f"2. Je schrijfrechten hebt in de map: {year_folder}\n"
                    f"3. De map niet alleen-lezen is\n"
                    f"Originele fout: {e}"
                )

    return filepath

def _create_header_box(ws, month_name, year):
    """Creëert een mooie header box bovenaan het document."""
    ws.merge_cells('B2:H3')
    title_cell = ws['B2']
    title_cell.value = f"WERKLOGBOEK - {month_name.upper()} {year}"
    title_cell.font = Font(bold=True, size=16, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    thick_border_side = Side(style='thick')
    
    for row in ws['B2:H3']:
        for cell in row:
            cell.border = Border(top=thick_border_side, bottom=thick_border_side, left=thick_border_side, right=thick_border_side)

    info_start_row = 5
    info_end_row = 7
    
    label_font = Font(bold=True)
    label_alignment = Alignment(horizontal="right", vertical="center")
    label_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    value_alignment = Alignment(horizontal="center", vertical="center")
    
    labels = {
        5: "Dagen gewerkt:",
        6: "Totaal uren aanwezig (bruto):",
        7: "Totaal uren gewerkt (netto):"
    }

    for row_num, text in labels.items():
        ws.merge_cells(f'B{row_num}:D{row_num}')
        label_cell = ws[f'B{row_num}']
        label_cell.value = text
        label_cell.font = label_font
        label_cell.alignment = label_alignment
        label_cell.fill = label_fill

        ws.merge_cells(f'E{row_num}:G{row_num}')
        ws[f'E{row_num}'].alignment = value_alignment

    # Randen voor de gehele info box
    thick_side = Side(style='thick')
    thin_side = Side(style='thin')
    for row_idx in range(info_start_row, info_end_row + 1):
        for col_idx in range(2, 9): # Moet tot 9 lopen voor kolom H
            cell = ws.cell(row=row_idx, column=col_idx)
            left = thin_side if col_idx > 2 else thick_side
            right = thin_side if col_idx < 8 else thick_side
            top = thin_side if row_idx > info_start_row else thick_side
            bottom = thin_side if row_idx < info_end_row else thick_side
            cell.border = Border(left=left, right=right, top=top, bottom=bottom)

    return info_end_row

def _setup_table_headers(ws, row):
    """Stelt de tabel headers in."""
    headers = ["Datum", "Begintijd", "Eindtijd", "Tijd Aanwezig", "Pauze (uren)", "Uren Gewerkt", "Activiteit"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thick_side = Side(style='thick')
    header_border = Border(top=thick_side, bottom=thick_side, left=thick_side, right=thick_side)

    for col_index, title in enumerate(headers, start=2):
        cell = ws.cell(row=row, column=col_index, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border

def _fill_weekday_data(ws, weekdays, start_row):
    """Vult de weekdag data in."""
    current_row = start_row
    current_week = None

    thin_side = Side(style='thin')
    data_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    week_separator_fill = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    dag_namen = ['Maandag', 'Dinsdag', 'Woensdag', 'Donderdag', 'Vrijdag']

    for date in weekdays:
        week_number = date.isocalendar()[1]

        if week_number != current_week:
            current_week = week_number
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=8)
            week_cell = ws.cell(row=current_row, column=2, value=f"Week {week_number}")
            week_cell.fill = week_separator_fill
            week_cell.font = Font(bold=True, color="FFFFFF")
            week_cell.alignment = Alignment(horizontal="center", vertical="center")
            week_cell.border = data_border
            current_row += 1

        dag_naam = dag_namen[date.weekday()]
        datum_text = f"{dag_naam} - {date.strftime('%d/%m')}"
        ws.cell(row=current_row, column=2, value=datum_text)
        
        # Voor kolom E: Tijd Aanwezig
        aanwezig_cell = ws.cell(row=current_row, column=5)
        aanwezig_cell.value = f'=IF(AND(ISNUMBER(C{current_row}), ISNUMBER(D{current_row})), (D{current_row}-C{current_row})*24, "")'
        aanwezig_cell.number_format = '0.##'

        # Voor kolom G: Uren Gewerkt
        gewerkt_cell = ws.cell(row=current_row, column=7)
        gewerkt_cell.value = (
            f'=IF(E{current_row}<>"", E{current_row}-N(F{current_row}), '
            f'IF(H{current_row}<>"", 8, ""))'
        )
        gewerkt_cell.number_format = '0.##'

        for col in range(2, 9):
            cell = ws.cell(row=current_row, column=col)
            cell.border = data_border
            
            if col == 8:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

            if (date.weekday() % 2) != 0:
                cell.fill = zebra_fill
        
        ws.cell(row=current_row, column=3).number_format = 'h:mm'
        ws.cell(row=current_row, column=4).number_format = 'h:mm'
        
        # Voor kolom F: Pauze (uren)
        ws.cell(row=current_row, column=6).number_format = '0.##'

        current_row += 1

    return current_row - 1

def _update_summary_formulas(ws, data_start_row, last_data_row):
    """Update de samenvattingsformules in de header."""
    if last_data_row < data_start_row:
        return
        
    ws["E5"].value = f'=COUNTIF(H{data_start_row}:H{last_data_row},"<>")'
    ws["E5"].number_format = '0'
    
    # Totaal uren aanwezig (bruto)
    ws["E6"].value = f'=SUM(E{data_start_row}:E{last_data_row})'
    ws["E6"].number_format = '0.##'

    # Totaal uren gewerkt (netto)
    ws["E7"].value = f'=SUM(G{data_start_row}:G{last_data_row})'
    ws["E7"].number_format = '0.##'

def _setup_column_widths(ws):
    """Stelt de kolom breedtes in."""
    column_widths = {'A': 2, 'B': 22, 'C': 12, 'D': 12, 'E': 15, 'F': 15, 'G': 15, 'H': 80}
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

def main():
    """Hoofdfunctie voor het testen van de logboek generator."""
    try:

        script_dir = os.path.dirname(os.path.abspath(__file__))
        print(f"📂 Script-locatie gedetecteerd: {script_dir}")
        print(f"   Output wordt opgeslagen relatief aan deze map.")

        print("🔐 Controleer schrijfrechten...")
        test_file = os.path.join(script_dir, "test_write.tmp")
        try:
            with open(test_file, 'w') as f:
                f.write("test")
            os.remove(test_file)
            print("✅ Schrijfrechten OK")
        except Exception as e:
            print(f"❌ Fout bij controleren van schrijfrechten: {e}")
            return

        print("📝 Logboek wordt aangemaakt...")
        # Voorbeeld: maak een logboek voor de huidige maand en jaar
        nu = datetime.now()
        maand_naam_nl = list(NEDERLANDS_MAANDEN.keys())[nu.month - 1]
        #filepath = create_logbook_excel(maand_naam_nl, nu.year, base_folder=script_dir)
        
        # Of een specifieke maand:
        filepath = create_logbook_excel("september", 2025, base_folder=script_dir)
        
        print(f"✅ Excel-bestand succesvol aangemaakt!")
        print(f"📄 Pad: {os.path.abspath(filepath)}")

    except PermissionError as e:
        print(f"🔒 Toegangsfout: {e}")
    except Exception as e:
        print(f"❌ Een onverwachte fout is opgetreden: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()